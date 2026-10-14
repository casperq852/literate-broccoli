"""Tests for the Word report builder: the exported document carries what the review says, edge cases don't crash."""

from __future__ import annotations

from io import BytesIO

from docx import Document

from dcf.dcf_parsing_model import ParsedRecord
from kpi_monitor.kpi_export_model import build_report
from kpi_monitor.kpi_model import (
    KpiAnalysis,
    KpiFinding,
    KpiReview,
    SourceFile,
    build_series,
    format_value,
    report_filename,
    series_decimals,
)
from search.search_model import Chunk

PERIODS = ["2024-06-30", "2024-09-30", "2024-12-31", "2025-03-31"]


def _rec(metric: str, period: str, value: float, *, est: bool = False, doc: str | None = None) -> ParsedRecord:
    return ParsedRecord(metric, None, f"{period}T00:00:00Z", est, value, "EUR", "Millions", doc)


def _analysis(
    findings: tuple[KpiFinding, ...], *, focus: str = "inventory", summary: str = "Margins fell."
) -> KpiAnalysis:
    records = [
        _rec("Gross Margin %", p, v, doc="2025-05-01") for p, v in zip(PERIODS, [51.5, 50.25, 49.1, 47.9], strict=True)
    ]
    records.append(
        _rec("Gross Margin %", "2025-03-31", 49.5, est=True, doc="2025-01-01")
    )  # earlier estimate of the latest period
    records.append(_rec("Gross Margin %", "2025-06-30", 48.0, est=True, doc="2025-05-01"))
    doc = Chunk("CFO: margin hit by mix.", None, None, None, None, "earnings_transcript", "2025-04-30", "Co", 1.0)
    return KpiAnalysis(
        review=KpiReview(summary, findings),
        as_of="2025-03-31",
        series=build_series(records),
        documents=[doc],
        focus_notes=focus,
        sources=(SourceFile("MODL_ACME.xlsx", "financial_model", "p/MODL_ACME.xlsx", 10),),
        created_at="2026-10-14T08:55:01+00:00",
    )


def _finding(metric: str = "Gross Margin %", assessment: str = "concerning") -> KpiFinding:
    return KpiFinding(
        metric, assessment, "Margin keeps falling", "Down 3.6pp over three quarters.", "51.5 -> 47.9", (1,), "Why?"
    )


def _open(data: bytes) -> Document:
    return Document(BytesIO(data))


def _all_text(document: Document) -> str:
    parts = [p.text for p in document.paragraphs]
    parts += [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    return "\n".join(parts)


class TestBuildReport:
    def test_report_carries_title_summary_findings_and_sources(self) -> None:
        data = build_report(
            _analysis((_finding(),)), "Acme NV", "ACME", {1: "Acme · earnings transcript · 2025-04-30"}, "2026-10-14"
        )
        text = _all_text(_open(data))
        for expected in (
            "Acme NV",
            "KPI review · ACME",
            "2025-03-31",
            "Analyst focus:",
            "inventory",
            "Margins fell.",
            "Concerning",
            "Margin keeps falling",
            "Down 3.6pp over three quarters.",
            "Follow-up:",
            "[1] Acme · earnings transcript · 2025-04-30",
            "MODL_ACME.xlsx",
            "CFO: margin hit by mix.",
        ):
            assert expected in text, expected

    def test_figures_table_has_consistent_decimals_and_prior_estimate_row(self) -> None:
        document = _open(build_report(_analysis((_finding(),)), "Acme", "ACME", {1: "x"}, "2026-10-14"))
        text = _all_text(document)
        assert "Prior estimate" in text and "49.50" in text  # earlier estimate of the latest reported period
        assert "51.50" in text and "50.25" in text  # every value in the series shows two decimals

    def test_each_finding_with_data_gets_a_chart(self) -> None:
        analysis = _analysis((_finding(), _finding(assessment="watch")))
        assert len(_open(build_report(analysis, "Acme", "ACME", {1: "x"}, "d")).inline_shapes) == 2
        assert len(_open(build_report(analysis, "Acme", "ACME", {1: "x"}, "d", charts=False)).inline_shapes) == 0

    def test_finding_for_a_metric_missing_from_the_model_still_exports(self) -> None:
        text = _all_text(
            _open(build_report(_analysis((_finding("No Such Metric", "neutral"),)), "Acme", "ACME", {1: "x"}, "d"))
        )
        assert "not found in the model table" in text

    def test_review_without_findings_or_summary_exports(self) -> None:
        document = _open(build_report(_analysis((), focus="", summary=""), "Acme", "ACME", {}, "2026-10-14"))
        text = _all_text(document)
        assert "Analyst focus" not in text and "Findings" not in text and "Acme" in text

    def test_page_numbers_are_fields_in_the_footer(self) -> None:
        document = _open(build_report(_analysis((_finding(),)), "Acme", "ACME", {1: "x"}, "d"))
        footer_xml = document.sections[0].footer._element.xml
        assert "PAGE" in footer_xml and "NUMPAGES" in footer_xml


class TestFormatting:
    def test_whole_numbers_and_large_values_have_no_decimals(self) -> None:
        assert format_value(700.0) == "700"
        assert format_value(1234.56, 2) == "1,235"
        assert format_value(12.5, 2) == "12.50"

    def test_series_decimals_is_two_if_any_value_is_fractional(self) -> None:
        whole = build_series([_rec("Rev", p, float(i * 100), doc="d") for i, p in enumerate(PERIODS, start=1)])[0]
        fractional = build_series([_rec("Rev", p, i + 0.5, doc="d") for i, p in enumerate(PERIODS, start=1)])[0]
        assert series_decimals(whole) == 0 and series_decimals(fractional) == 2

    def test_report_filename(self) -> None:
        assert report_filename("NOVO-B DC", "2025-03-31", "docx") == "KPI_review_novo-b-dc_2025-03-31.docx"
