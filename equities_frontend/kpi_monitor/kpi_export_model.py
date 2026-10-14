"""
Pure Word (.docx) report builder for the KPI Monitor's export button.

Turns a finished (or reopened) review into a clean, shareable document: a
summary, an at-a-glance table, then one section per finding with the analyst
commentary, a trend chart and the underlying figures, followed by the sources.
Everything is built from the ``KpiAnalysis`` passed in, so a saved review
exports identically to a fresh one. No network or Streamlit access.
"""

from __future__ import annotations

from datetime import date
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor, Twips
from matplotlib.figure import Figure

from kpi_monitor.kpi_model import (
    KpiAnalysis,
    KpiFinding,
    KpiSeries,
    format_value,
    series_decimals,
    series_for_metric,
)

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

LOGO_PATH = Path(__file__).resolve().parent.parent / "assets" / "asr_vermogensbeheer_logo.png"

# assessment -> (label, text colour, fill colour). The label always accompanies the colour.
_ASSESSMENT_STYLE: dict[str, tuple[str, str, str]] = {
    "concerning": ("Concerning", "B3261E", "FDECEA"),
    "watch": ("Watch", "8A5A00", "FFF1CC"),
    "positive": ("Positive", "1E6B31", "E3F3E7"),
    "neutral": ("Neutral", "4B5563", "EEF0F2"),
}
_INK = "1F2937"
_MUTED = "6B7280"
_RULE = "D1D5DB"
_HEADER_FILL = "E5E9EF"
_TABLE_PERIODS = 8  # most recent periods shown in each finding's figures table
_MAX_EXCERPT = 280


# ── low-level docx helpers ──────────────────────────────────────────────────────


def _rgb(hex_colour: str) -> RGBColor:
    return RGBColor.from_string(hex_colour)


def _shade(cell, fill: str) -> None:
    props = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)
    props.append(shading)


def _light_borders(table, colour: str = _RULE) -> None:
    props = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), colour)
        borders.append(element)
    props.append(borders)


def _fixed_widths(table, widths: list) -> None:
    """Pin column widths (Word and LibreOffice otherwise spread columns evenly)."""
    table.autofit = False
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    table._tbl.tblPr.append(layout)
    for column, width in zip(table.columns, widths, strict=True):
        column.width = width
    for row in table.rows:
        for cell, width in zip(row.cells, widths, strict=True):
            cell.width = width


def _keep_together(table) -> None:
    """Keep a table's rows on one page: no row splits, and each row stays with the next."""
    for index, row in enumerate(table.rows):
        trpr = row._tr.get_or_add_trPr()
        trpr.append(OxmlElement("w:cantSplit"))
        if index < len(table.rows) - 1:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.keep_with_next = True


def _cell_margins(table, top: int = 50, bottom: int = 50, side: int = 90) -> None:
    margins = OxmlElement("w:tblCellMar")
    for edge, value in (("top", top), ("left", side), ("bottom", bottom), ("right", side)):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")
        margins.append(element)
    table._tbl.tblPr.append(margins)


def _write(cell_or_par, text: str, *, bold: bool = False, size: float = 9, colour: str = _INK, align=None) -> None:
    paragraph = cell_or_par.paragraphs[0] if hasattr(cell_or_par, "paragraphs") else cell_or_par
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.color.rgb = _rgb(colour)
    if align is not None:
        paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(0)


def _field(paragraph, instruction: str, size: float = 8) -> None:
    """Insert a Word field (PAGE / NUMPAGES); each part is its own run so the text size applies to the result."""
    parts: list[tuple[str, str]] = [("fldChar", "begin"), ("instrText", instruction), ("fldChar", "separate"),
                                    ("t", "1"), ("fldChar", "end")]  # fmt: skip
    for tag, value in parts:
        run = paragraph.add_run()
        run.font.size = Pt(size)
        run.font.color.rgb = _rgb(_MUTED)
        element = OxmlElement(f"w:{tag}")
        if tag == "fldChar":
            element.set(qn("w:fldCharType"), value)
        else:
            element.text = f" {value} " if tag == "instrText" else value
            if tag == "instrText":
                element.set(qn("xml:space"), "preserve")
        run._r.append(element)


def _paragraph(document, text: str = "", *, size: float = 10, bold: bool = False, colour: str = _INK, after: float = 4):
    paragraph = document.add_paragraph()
    if text:
        run = paragraph.add_run(text)
        run.bold = bold
        run.font.size = Pt(size)
        run.font.color.rgb = _rgb(colour)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.space_before = Pt(0)
    return paragraph


def _heading(document, text: str, *, level: int) -> None:
    paragraph = document.add_heading(level=level)
    run = paragraph.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt({1: 15, 2: 12}[level])
    run.font.color.rgb = _rgb(_INK)
    paragraph.paragraph_format.space_before = Pt(14 if level == 1 else 12)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.keep_with_next = True


# ── page furniture ──────────────────────────────────────────────────────────────


def _right_tab_only(paragraph, width) -> None:
    """Replace the Header/Footer style's centre and right tab stops with a single right stop at the margin."""
    stops = paragraph.paragraph_format.tab_stops
    for position in (4680, 9360):  # the default template's centre and right stops, in twips
        stops.add_tab_stop(Twips(position), WD_TAB_ALIGNMENT.CLEAR)
    stops.add_tab_stop(width, WD_TAB_ALIGNMENT.RIGHT)


def _setup_page(document, ticker: str) -> None:
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.left_margin = section.right_margin = Cm(2)
    section.top_margin, section.bottom_margin = Cm(3.0), Cm(2)
    section.header_distance, section.footer_distance = Cm(0.9), Cm(0.9)
    width = section.page_width - section.left_margin - section.right_margin

    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)

    header = section.header.paragraphs[0]
    _right_tab_only(header, width)
    if LOGO_PATH.exists():
        header.add_run().add_picture(str(LOGO_PATH), height=Cm(1.7))
    run = header.add_run(f"\tEquities · KPI review · {ticker}")
    run.font.size = Pt(8)
    run.font.color.rgb = _rgb(_MUTED)

    footer = section.footer.paragraphs[0]
    _right_tab_only(footer, width)
    run = footer.add_run("Internal use · AI-assisted analysis: verify figures against the source model\tPage ")
    run.font.size = Pt(8)
    run.font.color.rgb = _rgb(_MUTED)
    _field(footer, "PAGE")
    run = footer.add_run(" of ")
    run.font.size = Pt(8)
    run.font.color.rgb = _rgb(_MUTED)
    _field(footer, "NUMPAGES")


# ── charts ──────────────────────────────────────────────────────────────────────


def _short_period(period_end: str) -> str:
    return date.fromisoformat(period_end).strftime("%b %y")


def _axis_label(value: float, _position: float) -> str:
    return f"{value:,.0f}" if abs(value) >= 100 or value == round(value) else f"{value:,.1f}"  # noqa: PLR2004


def _trend_chart(series: KpiSeries) -> bytes:
    """A small trend chart: reported values solid, estimates dashed with hollow markers."""
    points = series.points
    xs = list(range(len(points)))
    figure = Figure(figsize=(6.6, 1.9), dpi=200)
    axes = figure.subplots()

    reported = [(x, p.value) for x, p in zip(xs, points, strict=True) if not p.is_estimate]
    estimates = [(x, p.value) for x, p in zip(xs, points, strict=True) if p.is_estimate]
    if reported:
        axes.plot(
            *zip(*reported, strict=True), color=f"#{_INK}", linewidth=1.6, marker="o", markersize=3.5, label="Reported"
        )
        axes.plot(
            *reported[-1],
            marker="o",
            markersize=6.5,
            color="#2563EB",
            zorder=5,
            label="Latest reported",
            linestyle="none",
        )
    if estimates:
        link = ([reported[-1]] if reported else []) + estimates
        axes.plot(*zip(*link, strict=True), color="#6B7280", linewidth=1.3, linestyle="--", label="Estimate")
        axes.plot(*zip(*estimates, strict=True), linestyle="none", marker="o", markersize=4, markerfacecolor="white",
                  markeredgecolor="#6B7280")  # fmt: skip
    priors = [(x, p.prior_estimate) for x, p in zip(xs, points, strict=True) if p.prior_estimate is not None]
    if priors:
        axes.plot(*zip(*priors, strict=True), marker="_", markersize=9, color="#B3261E", linestyle="none", zorder=4,
                  label="Prior estimate")  # fmt: skip

    step = max(1, len(points) // 8)
    axes.set_xticks(xs[::step], [_short_period(p.period_end) for p in points][::step])
    axes.tick_params(labelsize=7, colors=f"#{_MUTED}", length=2)
    axes.yaxis.set_major_formatter(_axis_label)
    axes.grid(axis="y", color=f"#{_RULE}", linewidth=0.5)
    axes.set_axisbelow(True)
    for side in ("top", "right"):
        axes.spines[side].set_visible(False)
    for side in ("left", "bottom"):
        axes.spines[side].set_color(f"#{_RULE}")
    axes.margins(x=0.04, y=0.2)
    axes.legend(loc="lower right", bbox_to_anchor=(1.0, 1.0), ncol=4, fontsize=6.5, frameon=False, handlelength=1.6,
                borderaxespad=0.2, labelcolor=f"#{_MUTED}")  # fmt: skip
    figure.tight_layout(pad=0.4)

    buffer = BytesIO()
    figure.savefig(buffer, format="png")
    return buffer.getvalue()


# ── sections ────────────────────────────────────────────────────────────────────


def _title_block(document, analysis: KpiAnalysis, company: str, ticker: str, generated_on: str) -> None:
    title = _paragraph(document, company, size=22, bold=True, after=0)
    title.paragraph_format.space_before = Pt(4)
    _paragraph(document, f"KPI review · {ticker}", size=12, colour=_MUTED, after=6)

    table = document.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    _light_borders(table)
    _cell_margins(table)
    details = (
        ("Latest reported period", analysis.as_of),
        ("Review made", (analysis.created_at[:10] or generated_on)),
        ("Basis", f"{len(analysis.series)} metrics · {len(analysis.documents)} document passages"),
    )
    for cell, (label, value) in zip(table.rows[0].cells, details, strict=True):
        _shade(cell, "F7F8FA")
        _write(cell, label, size=7.5, colour=_MUTED)
        value_par = cell.add_paragraph()
        _write(value_par, value, bold=True, size=9.5)
    _fixed_widths(table, [Cm(4.2), Cm(4.2), Cm(8.6)])
    if analysis.focus_notes:
        _paragraph(document, "", after=2)
        focus = _paragraph(document, after=2)
        _write(focus, "Analyst focus: ", bold=True, size=9, colour=_MUTED)
        _write(focus, analysis.focus_notes, size=9, colour=_MUTED)


def _summary_block(document, analysis: KpiAnalysis) -> None:
    findings = analysis.review.findings
    if analysis.review.summary:
        _heading(document, "Summary", level=1)
        table = document.add_table(rows=1, cols=1)
        _light_borders(table)
        _cell_margins(table, top=110, bottom=110, side=150)
        cell = table.rows[0].cells[0]
        _shade(cell, "F7F8FA")
        _write(cell, analysis.review.summary, size=10.5)
        cell.paragraphs[0].paragraph_format.line_spacing = 1.15
        _fixed_widths(table, [Cm(17.0)])
    if not findings:
        return

    _heading(document, "At a glance", level=1)
    table = document.add_table(rows=1, cols=3)
    _light_borders(table)
    _cell_margins(table)
    for cell, text, width in zip(
        table.rows[0].cells, ("Assessment", "Metric", "Headline"), (Cm(2.6), Cm(6.4), Cm(8)), strict=True
    ):
        _shade(cell, _HEADER_FILL)
        _write(cell, text, bold=True, size=8.5)
        cell.width = width
    for finding in findings:
        label, colour, fill = _ASSESSMENT_STYLE[finding.assessment]
        row = table.add_row().cells
        _shade(row[0], fill)
        _write(row[0], label, bold=True, size=8.5, colour=colour)
        _write(row[1], finding.metric, size=8.5)
        _write(row[2], finding.headline or "", size=8.5)
    _fixed_widths(table, [Cm(2.6), Cm(6.4), Cm(8.0)])


def _figures_table(document, series: KpiSeries) -> None:
    points = series.points[-_TABLE_PERIODS:]
    has_prior = any(p.prior_estimate is not None for p in points)
    decimals = series_decimals(series)
    rows = 3 + has_prior
    table = document.add_table(rows=rows, cols=len(points) + 1)
    _light_borders(table)
    _cell_margins(table, top=25, bottom=25, side=60)
    labels = ["Period end", "Basis", "Value"] + (["Prior estimate"] if has_prior else [])
    for index, label in enumerate(labels):
        _shade(table.rows[index].cells[0], _HEADER_FILL)
        _write(table.rows[index].cells[0], label, bold=True, size=7.5)
    for col, point in enumerate(points, start=1):
        cells = [table.rows[i].cells[col] for i in range(rows)]
        _shade(cells[0], _HEADER_FILL)
        _write(cells[0], _short_period(point.period_end), bold=True, size=7.5, align=WD_ALIGN_PARAGRAPH.RIGHT)
        _write(
            cells[1],
            "Estimate" if point.is_estimate else "Reported",
            size=7.5,
            colour=_MUTED,
            align=WD_ALIGN_PARAGRAPH.RIGHT,
        )
        _write(cells[2], format_value(point.value, decimals), bold=True, size=8, align=WD_ALIGN_PARAGRAPH.RIGHT)
        if has_prior:
            prior = "" if point.prior_estimate is None else format_value(point.prior_estimate, decimals)
            _write(cells[3], prior, size=7.5, colour=_MUTED, align=WD_ALIGN_PARAGRAPH.RIGHT)
    first = Cm(2.6)
    _fixed_widths(table, [first] + [(Cm(17.0) - first) // len(points)] * len(points))
    _keep_together(table)
    unit = " ".join(part for part in (series.currency, series.unit) if part)
    if unit:
        _paragraph(document, f"Values in {unit}.", size=7.5, colour=_MUTED, after=2).paragraph_format.space_before = Pt(2)


def _finding(document, finding: KpiFinding, series: KpiSeries | None, sources: dict[int, str], *, charts: bool) -> None:
    label, colour, fill = _ASSESSMENT_STYLE[finding.assessment]
    _heading(document, finding.headline or finding.metric, level=2)

    tag = document.add_table(rows=1, cols=2)
    _cell_margins(tag, top=30, bottom=30, side=90)
    _shade(tag.rows[0].cells[0], fill)
    _write(tag.rows[0].cells[0], label.upper(), bold=True, size=8, colour=colour)
    _write(tag.rows[0].cells[1], f"  {finding.metric}", size=8.5, colour=_MUTED)
    _fixed_widths(tag, [Cm(2.6), Cm(14.4)])
    for cell in tag.rows[0].cells:
        cell.paragraphs[0].paragraph_format.keep_with_next = True

    body = _paragraph(document, finding.commentary, after=4)
    body.paragraph_format.keep_with_next = True
    body.paragraph_format.space_before = Pt(6)
    body.paragraph_format.line_spacing = 1.1
    for caption, text in (("Evidence", finding.evidence), ("Follow-up", finding.follow_up)):
        if text:
            line = _paragraph(document, after=3)
            line.paragraph_format.keep_with_next = True
            _write(line, f"{caption}: ", bold=True, size=9.5)
            _write(line, text, size=9.5)
    if finding.document_refs:
        refs = _paragraph(document, after=4)
        refs.paragraph_format.keep_with_next = True
        _write(refs, "Sources: ", bold=True, size=8.5, colour=_MUTED)
        _write(
            refs,
            "; ".join(f"[{ref}] {sources[ref]}" for ref in finding.document_refs if ref in sources),
            size=8.5,
            colour=_MUTED,
        )

    if series is None:
        _paragraph(
            document,
            "This metric name was not found in the model table; check the commentary against the model.",
            size=8.5,
            colour=_MUTED,
        )
        return
    if charts and len(series.points) > 1:
        picture = document.add_paragraph()
        picture.paragraph_format.keep_with_next = True
        picture.paragraph_format.space_after = Pt(2)
        picture.add_run().add_picture(BytesIO(_trend_chart(series)), width=Cm(16.4))
    _figures_table(document, series)


def _sources_section(document, analysis: KpiAnalysis, source_labels: dict[int, str]) -> None:
    if analysis.sources:
        _heading(document, "Files added for this review", level=1)
        for source in analysis.sources:
            kind = "Bloomberg model" if source.kind == "financial_model" else "Document"
            bullet = document.add_paragraph(style="List Bullet")
            _write(bullet, f"{source.filename} ", bold=True, size=9)
            _write(bullet, f"· {kind}", size=9, colour=_MUTED)
    if analysis.documents:
        _heading(document, "Document passages the review drew on", level=1)
        for index, chunk in enumerate(analysis.documents, start=1):
            line = _paragraph(document, after=1)
            line.paragraph_format.keep_with_next = True
            _write(line, f"[{index}] {source_labels.get(index, '')}", bold=True, size=8.5)
            excerpt = " ".join(chunk.content.split())
            excerpt = excerpt[:_MAX_EXCERPT].rstrip() + ("…" if len(excerpt) > _MAX_EXCERPT else "")
            _paragraph(document, excerpt, size=8.5, colour=_MUTED, after=5)


# ── entry point ─────────────────────────────────────────────────────────────────


def build_report(
    analysis: KpiAnalysis,
    company: str,
    ticker: str,
    source_labels: dict[int, str],
    generated_on: str,
    *,
    charts: bool = True,
) -> bytes:
    """
    Build the Word report for a KPI review.

    Args:
        analysis (KpiAnalysis): The review with its figures, passages and files.
        company (str): The company's display name.
        ticker (str): The company ticker.
        source_labels (dict[int, str]): Citation label per numbered document passage.
        generated_on (str): Date shown when the review itself carries no timestamp.
        charts (bool): Draw a trend chart per finding (set False to skip matplotlib).

    Returns:
        bytes: The ``.docx`` file.
    """
    document = Document()
    document.core_properties.title = f"KPI review: {company} ({ticker})"
    document.core_properties.subject = f"Latest reported period {analysis.as_of}"
    document.core_properties.author = "Equities Research RAG"
    _setup_page(document, ticker)

    _title_block(document, analysis, company, ticker, generated_on)
    _summary_block(document, analysis)
    if analysis.review.findings:
        _heading(document, "Findings", level=1)
    for finding in analysis.review.findings:
        _finding(document, finding, series_for_metric(analysis.series, finding.metric), source_labels, charts=charts)
    _sources_section(document, analysis, source_labels)

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()
