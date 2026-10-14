"""Pure Excel workbook builder for the DCF Model Builder's export button."""

from __future__ import annotations

from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.worksheet import Worksheet

from dcf.dcf_parsing_model import ParsedRecord
from dcf.dcf_valuation_model import DcfResult

_HEADER_FONT = Font(bold=True)


def _write_header_row(sheet: Worksheet, values: list[str]) -> None:
    """Append one bold header row to ``sheet``."""
    sheet.append(values)
    for cell in sheet[sheet.max_row]:
        cell.font = _HEADER_FONT


def _build_assumptions_sheet(
    sheet: Worksheet,
    company: str,
    ticker: str,
    assumptions: dict[str, dict[str, Any]],
    rationale: str,
) -> None:
    """Populate the Assumptions sheet: every editable input plus its source and any LLM rationale."""
    sheet.title = "Assumptions"
    sheet.append(["Company", company])
    sheet.append(["Ticker", ticker])
    sheet.append([])
    _write_header_row(sheet, ["Assumption", "Value", "Source"])
    for name, entry in assumptions.items():
        sheet.append([name, entry.get("value"), entry.get("source", "")])
    sheet.column_dimensions["A"].width = 28
    sheet.column_dimensions["B"].width = 14
    sheet.column_dimensions["C"].width = 32

    if rationale:
        sheet.append([])
        sheet.append(["LLM suggestion rationale"])
        sheet[sheet.max_row][0].font = _HEADER_FONT
        sheet.append([rationale])
        sheet[sheet.max_row][0].alignment = Alignment(wrap_text=True)


def _build_historicals_sheet(
    sheet: Worksheet,
    historicals: dict[str, list[ParsedRecord]],
    line_item_labels: dict[str, str],
    mapping: dict[str, str | None],
) -> None:
    """Populate the Historicals sheet: one row per mapped line item, one column per period."""
    sheet.title = "Historicals"

    all_periods = sorted({record.period_end_date[:10] for series in historicals.values() for record in series})
    _write_header_row(sheet, ["Line item", "Source metric (Bloomberg breadcrumb)", *all_periods])

    estimate_flags = ["Reported/Estimate", ""]
    period_estimate: dict[str, str] = {}
    for series in historicals.values():
        for record in series:
            period_estimate[record.period_end_date[:10]] = "Estimate" if record.is_estimate else "Reported"
    estimate_flags.extend(period_estimate.get(period, "") for period in all_periods)
    sheet.append(estimate_flags)

    for item_id, series in historicals.items():
        if not series:
            continue
        values_by_period = {record.period_end_date[:10]: record.value for record in series}
        label = line_item_labels.get(item_id, item_id)
        source_metric = mapping.get(item_id) or ""
        sheet.append([label, source_metric, *(values_by_period.get(period) for period in all_periods)])

    sheet.column_dimensions["A"].width = 28
    sheet.column_dimensions["B"].width = 40


def _build_projections_sheet(sheet: Worksheet, result: DcfResult, net_debt: float, shares_outstanding: float) -> None:
    """Populate the Projections sheet: year-by-year explicit projection plus the valuation summary."""
    sheet.title = "Projections"
    years = [f"Year {i + 1}" for i in range(len(result.projected_revenue))]

    _write_header_row(sheet, ["", *years])
    sheet.append(["Revenue", *result.projected_revenue])
    sheet.append(["EBIT", *result.projected_ebit])
    sheet.append(["NOPAT", *result.projected_nopat])
    sheet.append(["Unlevered FCF", *result.projected_fcf])
    sheet.append(["Discount factor", *result.discount_factors])
    sheet.append(["PV of FCF", *result.pv_fcf])

    sheet.append([])
    _write_header_row(sheet, ["Valuation summary", "Value"])
    sheet.append(["Terminal value", result.terminal_value])
    sheet.append(["PV of terminal value", result.pv_terminal_value])
    sheet.append(["Enterprise value", result.enterprise_value])
    sheet.append(["Net debt", net_debt])
    sheet.append(["Equity value", result.equity_value])
    sheet.append(["Shares outstanding", shares_outstanding])
    sheet.append(["Value per share", result.value_per_share])

    sheet.column_dimensions["A"].width = 22


def build_workbook(
    company: str,
    ticker: str,
    mapping: dict[str, str | None],
    historicals: dict[str, list[ParsedRecord]],
    line_item_labels: dict[str, str],
    assumptions: dict[str, dict[str, Any]],
    result: DcfResult,
    net_debt: float,
    shares_outstanding: float,
    rationale: str = "",
) -> bytes:
    """
    Build the DCF workbook: Assumptions, Historicals, and Projections sheets.

    Args:
        company (str): The company name, shown at the top of the Assumptions sheet.
        ticker (str): The company ticker, shown at the top of the Assumptions sheet.
        mapping (dict[str, str | None]): canonical_item_id -> chosen metric
            breadcrumb, shown as a "Source metric" column in Historicals so
            the mapping is auditable from the exported file alone.
        historicals (dict[str, list[ParsedRecord]]): canonical_item_id ->
            historical time series.
        line_item_labels (dict[str, str]): canonical_item_id -> display label.
        assumptions (dict[str, dict[str, Any]]): assumption name -> {"value",
            "source"} shown in the Assumptions sheet.
        result (DcfResult): The computed DCF valuation.
        net_debt (float): Net debt used in the valuation.
        shares_outstanding (float): Shares outstanding used in the valuation.
        rationale (str): The LLM suggestion's rationale text, if any.

    Returns:
        bytes: The workbook serialized via openpyxl, ready for
        ``st.download_button``.
    """
    workbook = Workbook()
    _build_assumptions_sheet(workbook.active, company, ticker, assumptions, rationale)
    _build_historicals_sheet(workbook.create_sheet(), historicals, line_item_labels, mapping)
    _build_projections_sheet(workbook.create_sheet(), result, net_debt, shares_outstanding)

    workbook.properties.title = f"{company} DCF"

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
