"""
Controller for Bloomberg financial-model extraction.

Wires the pure parsing functions in ``financial_model_parsing_model`` together
into one record per metric x period. Expected sheet layout::

    Row 1   Company name
    Row 2   "<TICKER> ... Periodicity:<Q|S|A> Currency:<CCY> ..."
    Row 3   "In Millions of <CCY>" then period labels from column data_start_col
    Row 4   period end dates (datetimes) from column data_start_col
    Row 5+  metric rows: col 0 name (indented), col 1 Bloomberg field, col 4+ values

``extract_financial_model`` is kept as a module-level function so existing
importers (``vectorization/__init__.py``,
``vectorization/extraction/extraction_controller.py``) do not need to change
their call sites; it is a thin compatibility wrapper around
``_MetricReader``, which owns the stateful section-hierarchy walk that a pure
model function cannot express cleanly (each row's meaning depends on the
section headers seen in prior rows).
"""

from __future__ import annotations

import logging
from datetime import datetime
from io import BytesIO
from typing import Any

from openpyxl import load_workbook

from vectorization.financial_model import financial_model_parsing_model as parsing_model
from vectorization.financial_model.types import MetricRow, ModelHeader, Period

logger = logging.getLogger(__name__)

# Columns 0-3 hold metadata (name, Bloomberg field, calc field, segment id);
# period values start here. Overridable via config.yaml's
# financial_model_extraction.data_start_col.
DEFAULT_DATA_START_COL = 4

# Metric names containing these tokens are ratios/percentages: rendered to 2dp
# with no currency/unit suffix.
RATIO_TOKENS: tuple[str, ...] = ("eps", "%", "margin", "ratio", "per share", "yield", "cover", "payout")


def extract_financial_model(
    data: bytes,
    filename: str,
    data_start_col: int = DEFAULT_DATA_START_COL,
) -> list[dict[str, Any]]:
    """
    Extract one record per metric x period from a Bloomberg model workbook.

    Each record carries a full breadcrumb ``metric`` path, the exact
    ``period_end_date``, an ``is_estimate`` flag, and ``latest_reported_date``
    (the most recent reported period, used downstream as the document date).

    Args:
        data (bytes): Raw bytes of the uploaded ``.xlsx`` file.
        filename (str): Original filename, used only for logging.
        data_start_col (int): Index of the first period-value column. Defaults
            to the standard Bloomberg export layout.

    Returns:
        list[dict[str, Any]]: One record per metric x period, or an empty list
        on failure or if no period columns are found (letting the caller fall
        back to generic extraction).
    """
    try:
        workbook = load_workbook(BytesIO(data), data_only=True, read_only=True)
    except Exception as exc:  # noqa: BLE001 - corrupt upload, let caller fall back
        logger.warning("Cannot open %s: %s", filename, exc)
        return []

    sheet = workbook[_find_data_sheet(workbook)]
    periods = parsing_model.parse_periods(sheet, data_start_col=data_start_col)
    if not periods:
        if _has_uncached_formulas(data, sheet.title, data_start_col):
            logger.warning(
                "No period columns detected in %s: row 3/4's period-label and date "
                "cells appear to be formulas with no cached value, not a layout "
                "mismatch (openpyxl only reads what was cached at save time, unlike "
                "Excel's own live formula engine). Common when a Bloomberg export "
                "was never re-saved by Excel after generation -- open the file in "
                "Excel, force a recalculation (Ctrl+Alt+F9), save, and re-upload.",
                filename,
            )
        else:
            logger.warning("No period columns detected in %s; caller should fall back", filename)
        return []

    header = parsing_model.parse_header(sheet)
    unit = parsing_model.parse_unit(sheet)
    records = _MetricReader(header=header, periods=periods, unit=unit).read(sheet)

    parsing_model.stamp_latest_reported_date(records)
    logger.info(
        "Extracted %d records from %s (latest reported: %s)",
        len(records),
        filename,
        records[0]["latest_reported_date"] if records else None,
    )
    return records


def _has_uncached_formulas(data: bytes, sheet_name: str, data_start_col: int) -> bool:
    """
    True if row 3/4's period-label/date columns hold formulas rather than
    literal values -- the likely cause when ``parse_periods`` finds nothing
    despite the file looking correct in Excel itself (data_only=True returns
    None for a formula cell with no cached result, which Excel's own live
    calculation engine papers over when you view the file interactively).

    Reopens the workbook with ``data_only=False`` (the first load already
    used ``data_only=True`` to read literal values, which cannot also see
    formula source text) -- an extra parse, but only on the failure path
    ``extract_financial_model`` is already about to return an empty result
    from, so it costs nothing on the common case.

    Args:
        data: Raw bytes of the workbook (same bytes already parsed once).
        sheet_name: The sheet ``parse_periods`` was run against.
        data_start_col: Index of the first period-value column.

    Returns:
        bool: Whether any row-3/row-4 data column holds a formula string.
    """
    try:
        formula_workbook = load_workbook(BytesIO(data), data_only=False, read_only=True)
        rows = list(formula_workbook[sheet_name].iter_rows(min_row=3, max_row=4, values_only=True))
    except Exception:  # noqa: BLE001 - best-effort diagnostic, never block the real fallback
        return False
    return any(isinstance(value, str) and value.startswith("=") for row in rows for value in row[data_start_col:])


def _find_data_sheet(workbook: Any) -> str:
    """
    Pick the sheet with the most datetimes in row 4 (the period date row).

    This is a dataset-specific heuristic (Bloomberg workbooks may include cover
    sheets/notes alongside the actual model), so it stays in the Controller
    rather than the parsing model.

    Args:
        workbook (Any): An openpyxl workbook opened in read-only mode.

    Returns:
        str: The name of the sheet most likely to hold the financial model.
    """
    best_name, best_count = workbook.sheetnames[0], 0
    for name in workbook.sheetnames:
        try:
            row4 = next(workbook[name].iter_rows(min_row=4, max_row=4, max_col=30, values_only=True), ())
        except Exception as exc:  # noqa: BLE001
            logger.debug("Skipping sheet %s: %s", name, exc)
            continue
        count = sum(1 for value in row4 if isinstance(value, datetime))
        if count > best_count:
            best_name, best_count = name, count
    return best_name


class _MetricReader:
    """
    Walks the data rows of a sheet, tracking the section hierarchy to build
    breadcrumb metric paths and emitting one record per metric x period.
    """

    def __init__(self, header: ModelHeader, periods: dict[int, Period], unit: str) -> None:
        self._header = header
        self._periods = periods
        self._unit = unit
        self._sections: list[tuple[int, str]] = []  # (indent, name) stack

    def read(self, sheet: Any) -> list[dict[str, Any]]:
        """Walk every data row from row 5 onward, returning all metric x period records."""
        records: list[dict[str, Any]] = []
        for row in sheet.iter_rows(min_row=5, values_only=True):
            if not row[0]:
                continue
            name = str(row[0]).strip()
            if not name:
                continue

            indent = parsing_model.indent_depth(str(row[0]))
            if parsing_model.is_row_section_header(row, self._periods):
                self._push_section(indent=indent, name=name)
                continue

            records.extend(self._records_for_row(row=row, name=name, indent=indent))
        return records

    # ── per-row helpers ───────────────────────────────────────────────────────

    def _records_for_row(self, row: tuple, name: str, indent: int) -> list[dict[str, Any]]:
        """Emit one record per period that has a numeric value in this row."""
        parents = [section for depth, section in self._sections if depth < indent]
        metric = MetricRow(
            path=parsing_model.build_metric_path(parents=parents, leaf=name),
            leaf_name=name,
            segment=parents[-1] if parents else None,
            bloomberg_field=str(row[1]).strip() if len(row) > 1 and row[1] else "",
            indent=indent,
        )
        return [
            self._build_record(value=row[col], metric=metric, period=period)
            for col, period in self._periods.items()
            if col < len(row) and isinstance(row[col], (int, float))
        ]

    def _build_record(self, value: float, metric: MetricRow, period: Period) -> dict[str, Any]:
        """Assemble one metric x period record, including its rendered content string."""
        value_str, unit_suffix = parsing_model.format_metric_value(
            value=value,
            metric_path=metric.path,
            currency=self._header.currency,
            unit=self._unit,
            ratio_tokens=RATIO_TOKENS,
        )
        flag = "Estimate" if period.is_estimate else "Reported"
        end_str = period.end_date.strftime("%Y-%m-%d")
        content = (
            f"{self._header.company_name} - {metric.path} - "
            f"{period.label} (period ending {end_str}) ({flag}): {value_str}{unit_suffix}"
        )
        return {
            "content": content,
            "metric": metric.path,
            "metric_name": metric.leaf_name,
            "bloomberg_field": metric.bloomberg_field,
            "period_label": period.label,
            "period_end_date": period.end_date.strftime("%Y-%m-%dT%H:%M:%SZ"),
            "period_end_str": end_str,
            "is_estimate": "estimate" if period.is_estimate else "reported",
            "periodicity": self._header.periodicity,
            "currency": self._header.currency,
            "unit": self._unit,
            "segment": metric.segment,
            "company_name": self._header.company_name,
            "ticker": self._header.ticker,
            "indent_depth": metric.indent,
            "latest_reported_date": None,  # filled after all rows are read
        }

    # ── section stack ─────────────────────────────────────────────────────────

    def _push_section(self, indent: int, name: str) -> None:
        """Add a section header to the stack, popping siblings and children first."""
        while self._sections and self._sections[-1][0] >= indent:
            self._sections.pop()
        self._sections.append((indent, name))
