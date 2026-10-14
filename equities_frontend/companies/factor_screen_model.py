"""
Pure functions for parsing a periodically-uploaded Bloomberg factor/crowding
screen export.

Inspired by the "Human + Machine + Network Intelligence" framing in a UBS
research note the user shared: crowding/positioning data is treated there as
a distinct signal for judging whether a thesis is already priced. This
module turns a raw screen export into plain rows for
``database.cache.upsert_factor_screen_rows``; see
``companies.company_controller.CompanyController.factor_screen_for_company``
for how a row gets matched onto one of this app's own companies.
"""

from __future__ import annotations

import re
from io import BytesIO

from openpyxl import load_workbook

# Maps a screen export's header text (lowercased, stripped) to this app's
# own field names -- only the columns this app actually uses; the file's
# ~20 other HOLT-specific columns (CFROI, eCAP Award, Probability of
# Default, etc.) are simply not in this map and so are dropped.
_COLUMN_NAME_MAP: dict[str, str] = {
    "name": "name",
    "symbol": "symbol",
    "country": "country",
    "bloomberg identifier": "bloomberg_identifier",
    "sedol": "sedol",
    "overall factor percentile [region sector size relative]": "overall_factor_pct",
    "operational quality factor percentile [region sector size relative]": "quality_factor_pct",
    "low volatility factor percentile [region sector size relative]": "low_vol_factor_pct",
    "growth factor percentile [region sector size relative]": "growth_factor_pct",
    "momentum factor percentile [region sector size relative]": "momentum_factor_pct",
    "valuation factor percentile [region sector size relative]": "valuation_factor_pct",
    "ubs quant research crowding factor": "crowding_factor",
}

_NUMERIC_FIELDS = frozenset(
    {
        "overall_factor_pct",
        "quality_factor_pct",
        "low_vol_factor_pct",
        "growth_factor_pct",
        "momentum_factor_pct",
        "valuation_factor_pct",
        "crowding_factor",
    }
)

# How many leading rows to scan for the header row -- confirmed at row 12 in
# the sample export, comfortably inside this margin.
_HEADER_SCAN_ROWS = 20

_MONTH_ABBREVIATIONS: dict[str, int] = {
    "jan": 1, "feb": 2, "mar": 3, "apr": 4, "may": 5, "jun": 6,
    "jul": 7, "aug": 8, "sep": 9, "oct": 10, "nov": 11, "dec": 12,
}  # fmt: skip

_FILENAME_DATE_RE = re.compile(r"_(\d{4})_([A-Za-z]{3})_(\d{1,2})\.xlsx$", re.IGNORECASE)


def parse_screen_date(filename: str) -> str:
    """
    Parse the as-of date out of a screen export's filename.

    Args:
        filename: e.g. ``"Momentum Screen_2026_Aug_24.xlsx"``.

    Returns:
        str: ``YYYY-MM-DD``.

    Raises:
        ValueError: If the filename doesn't end in
            ``"_<YYYY>_<Mon>_<DD>.xlsx"``, or the month abbreviation isn't
            recognized -- raised rather than guessing, since the date drives
            this upload's primary key.
    """
    match = _FILENAME_DATE_RE.search(filename)
    if not match:
        raise ValueError(
            f"Could not find a '_<YYYY>_<Mon>_<DD>.xlsx' date in {filename!r} -- expected the same "
            'naming convention as "Momentum Screen_2026_Aug_24.xlsx".'
        )
    year, month_abbr, day = match.groups()
    month = _MONTH_ABBREVIATIONS.get(month_abbr.lower())
    if month is None:
        raise ValueError(f"Unrecognized month abbreviation {month_abbr!r} in {filename!r}.")
    return f"{year}-{month:02d}-{int(day):02d}"


def _find_header(sheet) -> tuple[int, dict[int, str]] | None:  # noqa: ANN001 - openpyxl worksheet, no clean public type
    """Scan the sheet's first rows for the one with both 'Name' and 'Symbol' headers; return (row_number, {col_index: field})."""
    for row_number, row in enumerate(sheet.iter_rows(min_row=1, max_row=_HEADER_SCAN_ROWS, values_only=True), start=1):
        candidate = {
            index: _COLUMN_NAME_MAP[str(cell).strip().lower()]
            for index, cell in enumerate(row)
            if cell is not None and str(cell).strip().lower() in _COLUMN_NAME_MAP
        }
        if "name" in candidate.values() and "symbol" in candidate.values():
            return row_number, candidate
    return None


def parse_factor_screen(data: bytes) -> list[dict]:
    """
    Parse a Bloomberg factor/crowding screen export into plain row dicts.

    Locates the header row by content (found by scanning, not a hardcoded
    row number, in case a future export shifts it), maps only the columns
    this app uses by header text rather than column position, and iterates
    every row after it. Blank/whitespace-only cells (the file's own
    convention for "no value" -- confirmed on the crowding-factor column)
    become ``None``.

    Args:
        data: Raw bytes of the uploaded ``.xlsx`` file.

    Returns:
        list[dict]: One dict per company, keys matching
        ``database.cache``'s ``factor_screen`` columns (``name``,
        ``symbol``, ``country``, ..., ``crowding_factor``). Rows with no
        ``name`` value are skipped.

    Raises:
        ValueError: If no row in the first ``_HEADER_SCAN_ROWS`` rows
            contains both "Name" and "Symbol" headers -- this doesn't look
            like a Bloomberg factor screen export.
    """
    workbook = load_workbook(BytesIO(data), data_only=True, read_only=True)
    sheet = workbook[workbook.sheetnames[0]]

    header = _find_header(sheet)
    if header is None:
        raise ValueError(
            f"Could not find a header row with both 'Name' and 'Symbol' columns in the first "
            f"{_HEADER_SCAN_ROWS} rows -- this doesn't look like a Bloomberg factor screen export."
        )
    header_row_number, header_by_index = header

    rows: list[dict] = []
    for row in sheet.iter_rows(min_row=header_row_number + 1, values_only=True):
        record: dict = {}
        for index, field in header_by_index.items():
            value = row[index] if index < len(row) else None
            if isinstance(value, str) and not value.strip():
                value = None
            record[field] = value

        if not record.get("name"):
            continue

        for field in _NUMERIC_FIELDS:
            value = record.get(field)
            if value is not None and not isinstance(value, int | float):
                try:
                    value = float(value)
                except (TypeError, ValueError):
                    value = None
            record[field] = value

        rows.append(record)

    return rows
